# -*- coding: utf-8 -*-
"""DOCX building helpers: Chinese typography, editorial tables, callouts, TOC."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x2E, 0x5C, 0x8A)
INK = RGBColor(0x22, 0x30, 0x3F)
GRAY = RGBColor(0x5D, 0x6B, 0x7A)
LIGHTGRAY = RGBColor(0x8A, 0x94, 0xA6)
AMBER = RGBColor(0xC0, 0x80, 0x3A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
LATIN_BODY = "Times New Roman"
CONTENT_WIDTH_CM = 16.2


def _set_run_font(run, name, size=None, bold=None, color=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    latin = LATIN_BODY if name == BODY_FONT else name
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), latin)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _shade(el, hex_fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    el.insert_element_before(shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
                             "w:vAlign", "w:hideMark", "w:headers", "w:cellIns",
                             "w:cellDel", "w:cellMerge", "w:tcPrChange")


def _cell_shade(cell, hex_fill):
    _shade(cell._tc.get_or_add_tcPr(), hex_fill)


def _cell_margins(table, top=40, bottom=40, left=90, right=90):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.insert_element_before(mar, "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def _set_grid(table, widths_cm):
    twips = [int(round(w / 2.54 * 1440)) for w in widths_cm]
    tbl = table._tbl
    grid = tbl.tblGrid
    for gc in list(grid.gridCol_lst):
        grid.remove(gc)
    for w in twips:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.insert_element_before(tblW, "w:jc", "w:tblCellSpacing", "w:tblInd",
                                    "w:tblBorders", "w:shd", "w:tblLayout",
                                    "w:tblCellMar", "w:tblLook", "w:tblCaption",
                                    "w:tblDescription", "w:tblPrChange")
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(sum(twips)))


def _table_borders(table, color="D5DCE6"):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    specs = {
        "top": ("single", 10, "1F3A5F"),
        "bottom": ("single", 10, "1F3A5F"),
        "left": ("none", 0, "auto"),
        "right": ("none", 0, "auto"),
        "insideH": ("single", 4, color),
        "insideV": ("none", 0, "auto"),
    }
    for edge, (val, sz, col) in specs.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), col)
        borders.append(el)
    tblPr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar",
                                 "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.insert_element_before(el, "w:trHeight", "w:tblHeader", "w:tblCellSpacing",
                               "w:jc", "w:hidden", "w:ins", "w:del", "w:trPrChange")


def _cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.insert_element_before(el, "w:trHeight", "w:tblHeader", "w:tblCellSpacing",
                               "w:jc", "w:hidden", "w:ins", "w:del", "w:trPrChange")


def _para_border(paragraph, edge="bottom", color="1F3A5F", sz=8, space=4):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.insert_element_before(pBdr, "w:shd", "w:tabs", "w:suppressAutoHyphens",
                                  "w:spacing", "w:ind", "w:contextualSpacing", "w:jc",
                                  "w:textDirection", "w:textAlignment", "w:outlineLvl",
                                  "w:rPr", "w:sectPr", "w:pPrChange")
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pBdr.append(el)


def _keep_with_next(paragraph, on=True):
    paragraph.paragraph_format.keep_with_next = on


def _add_field(paragraph, instr, placeholder="1"):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = instr
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = placeholder
    fldChar3 = OxmlElement("w:fldChar")
    fldChar3.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(t)
    run._r.append(fldChar3)
    return run


def setup_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.45)
    sec.bottom_margin = Cm(2.20)
    sec.left_margin = Cm(2.40)
    sec.right_margin = Cm(2.40)
    sec.header_distance = Cm(1.30)
    sec.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), BODY_FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.6
    pf.space_before = Pt(0)
    pf.space_after = Pt(3)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Pt(21)

    # document defaults for compatibility
    settings = doc.settings.element
    return doc


# --------------------------------------------------------------------------- blocks
class Builder:
    def __init__(self, doc, page_map=None):
        self.doc = doc
        self.page_map = page_map or {}
        self.headings = []  # (level, number, title, anchor_text)
        self.fig_no = 0
        self.tab_no = 0
        self.first_h1_done = False

    # ---- basic paragraphs
    def body(self, text, indent=True, size=10.5, color=INK, space_after=3,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.6, bold=False):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = line
        pf.space_after = Pt(space_after)
        pf.alignment = align
        pf.first_line_indent = Pt(21) if indent else Pt(0)
        self._rich(p, text, size=size, color=color, bold=bold)
        return p

    def _rich(self, p, text, size=10.5, color=INK, bold=False):
        """Supports **bold** spans."""
        parts = text.split("**")
        for i, part in enumerate(parts):
            if not part:
                continue
            r = p.add_run(part)
            _set_run_font(r, BODY_FONT, size, bold or (i % 2 == 1), color)

    def lead(self, text):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = 1.55
        pf.space_before = Pt(2)
        pf.space_after = Pt(9)
        pf.left_indent = Cm(0.42)
        pf.first_line_indent = Pt(0)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        el = OxmlElement("w:left")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "18")
        el.set(qn("w:space"), "10"); el.set(qn("w:color"), "C0803A")
        pBdr.append(el)
        if pBdr.getparent() is None:
            pPr.insert_element_before(pBdr, "w:shd", "w:tabs", "w:spacing", "w:ind",
                                      "w:contextualSpacing", "w:jc", "w:rPr", "w:pPrChange")
        self._rich(p, text, size=10.5, color=RGBColor(0x3D, 0x4B, 0x5C))

    def bullets(self, items, size=10.0):
        for it in items:
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.95)
            pf.first_line_indent = Cm(-0.42)
            pf.line_spacing = 1.5
            pf.space_after = Pt(2)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run("▪ ")
            _set_run_font(r, HEAD_FONT, 8, True, AMBER)
            self._rich(p, it, size=size)

    def kv(self, items):
        for k, v in items:
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.42)
            pf.first_line_indent = Pt(0)
            pf.line_spacing = 1.55
            pf.space_after = Pt(2)
            r = p.add_run(k + "　")
            _set_run_font(r, HEAD_FONT, 9.5, True, NAVY)
            self._rich(p, v, size=10.0)

    # ---- headings
    def h1(self, number, title, pagebreak=True):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        if pagebreak and self.first_h1_done:
            pf.page_break_before = True
        self.first_h1_done = True
        pf.space_before = Pt(2)
        pf.space_after = Pt(4)
        pf.keep_with_next = True
        pf.first_line_indent = Pt(0)
        r = p.add_run(number)
        _set_run_font(r, HEAD_FONT, 10, True, AMBER)
        p2 = self.doc.add_paragraph()
        pf = p2.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(12)
        pf.keep_with_next = True
        pf.first_line_indent = Pt(0)
        _para_border(p2, "bottom", "C9D2DE", 6, 6)
        r = p2.add_run(title)
        _set_run_font(r, HEAD_FONT, 19, True, NAVY)
        self.headings.append((1, number.strip(), title.strip(), f"{number} {title}"))
        # bookmark anchor
        return p2

    def h2(self, number, title):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(12)
        pf.space_after = Pt(5)
        pf.keep_with_next = True
        pf.first_line_indent = Pt(0)
        r = p.add_run(number + "　")
        _set_run_font(r, HEAD_FONT, 12.5, True, AMBER)
        r = p.add_run(title)
        _set_run_font(r, HEAD_FONT, 13.5, True, BLUE)
        self.headings.append((2, number.strip(), title.strip(), f"{number} {title}"))

    def h3(self, number, title):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(8)
        pf.space_after = Pt(3)
        pf.keep_with_next = True
        pf.first_line_indent = Pt(0)
        r = p.add_run((number + "　") if number else "")
        _set_run_font(r, HEAD_FONT, 9.5, True, AMBER)
        r = p.add_run(title)
        _set_run_font(r, HEAD_FONT, 11, True, INK)

    def ref(self, idx, text):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.9)
        pf.first_line_indent = Cm(-0.9)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.space_after = Pt(1.6)
        pf.line_spacing = 1.28
        r = p.add_run(f"[{idx}]\u3000")
        _set_run_font(r, HEAD_FONT, 8.0, True, NAVY)
        r = p.add_run(text)
        _set_run_font(r, BODY_FONT, 8.1, False, INK)

    def stats(self, items):
        n = len(items)
        t = self.doc.add_table(rows=1, cols=n)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
            val = "single" if edge == "top" else "none"
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), val)
            el.set(qn("w:sz"), "6" if edge == "top" else "0")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "C9D2DE")
            borders.append(el)
        tblPr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar",
                                    "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")
        _cell_margins(t, top=150, bottom=60, left=110, right=80)
        for i, (num, label) in enumerate(items):
            cell = t.rows[0].cells[i]
            cell.width = Cm(CONTENT_WIDTH_CM / n)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(num)
            _set_run_font(r, HEAD_FONT, 17, True, NAVY)
            p2 = cell.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p2.paragraph_format.space_before = Pt(2)
            p2.paragraph_format.space_after = Pt(0)
            p2.paragraph_format.first_line_indent = Pt(0)
            p2.paragraph_format.line_spacing = 1.25
            r = p2.add_run(label)
            _set_run_font(r, BODY_FONT, 7.2, False, GRAY)
        self.spacer(10)

    def label(self, text):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(10)
        pf.space_after = Pt(4)
        pf.first_line_indent = Pt(0)
        pf.keep_with_next = True
        r = p.add_run(text)
        _set_run_font(r, HEAD_FONT, 11, True, NAVY)
        _para_border(p, "bottom", "C9D2DE", 4, 3)

    # ---- callout
    def callout(self, label, text, fill="F4F1EC", bar="C0803A"):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        cell = t.cell(0, 0)
        cell.width = Cm(CONTENT_WIDTH_CM)
        _cell_shade(cell, fill)
        _cell_margins(t, top=110, bottom=110, left=170, right=150)
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge, (val, sz, col) in {
            "top": ("none", 0, "auto"), "bottom": ("none", 0, "auto"),
            "right": ("none", 0, "auto"), "insideH": ("none", 0, "auto"),
            "insideV": ("none", 0, "auto"), "left": ("single", 18, bar),
        }.items():
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), val); el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0"); el.set(qn("w:color"), col)
            borders.append(el)
        tblPr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar",
                                 "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.first_line_indent = Pt(0)
        r = p.add_run(label + "　")
        _set_run_font(r, HEAD_FONT, 9, True, RGBColor(0x7A, 0x4E, 0x50))
        self._rich(p, text, size=9.2, color=RGBColor(0x3D, 0x4B, 0x5C))
        self.spacer(4)

    def spacer(self, pt=6):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run("")
        _set_run_font(r, BODY_FONT, pt)

    # ---- figure
    def figure(self, path, caption, width_cm=CONTENT_WIDTH_CM, note=None):
        self.fig_no += 1
        p = self.doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Pt(0)
        p.add_run().add_picture(path, width=Cm(width_cm))
        cap = self.doc.add_paragraph()
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(2)
        cap.paragraph_format.line_spacing = 1.3
        cap.paragraph_format.first_line_indent = Pt(0)
        r = cap.add_run(f"图 {self.fig_no}　")
        _set_run_font(r, HEAD_FONT, 8.8, True, NAVY)
        r = cap.add_run(caption)
        _set_run_font(r, BODY_FONT, 8.8, False, GRAY)
        if note:
            n = self.doc.add_paragraph()
            n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            n.paragraph_format.space_after = Pt(10)
            n.paragraph_format.line_spacing = 1.25
            n.paragraph_format.first_line_indent = Pt(0)
            r = n.add_run("　　" + note)
            _set_run_font(r, BODY_FONT, 7.6, False, LIGHTGRAY)
        else:
            self.spacer(8)

    # ---- table
    def table(self, caption, header, rows, widths, note=None, font=8.4,
              header_font=8.2, first_col_bold=True, align=None, zebra=True):
        self.tab_no += 1
        cap = self.doc.add_paragraph()
        cap.paragraph_format.space_before = Pt(8)
        cap.paragraph_format.space_after = Pt(3)
        cap.paragraph_format.line_spacing = 1.3
        cap.paragraph_format.first_line_indent = Pt(0)
        cap.paragraph_format.keep_with_next = True
        r = cap.add_run(f"表 {self.tab_no}　")
        _set_run_font(r, HEAD_FONT, 8.8, True, NAVY)
        r = cap.add_run(caption)
        _set_run_font(r, BODY_FONT, 8.8, True, RGBColor(0x3D, 0x4B, 0x5C))

        ncols = len(header)
        t = self.doc.add_table(rows=1, cols=ncols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        _table_borders(t)
        _set_grid(t, widths)
        _cell_margins(t, top=55, bottom=55, left=80, right=80)

        hdr = t.rows[0]
        _repeat_header(hdr)
        _cant_split(hdr)
        for i, txt in enumerate(header):
            cell = hdr.cells[i]
            cell.width = Cm(widths[i])
            _cell_shade(cell, "1F3A5F")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.first_line_indent = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if widths[i] < 1.9 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(txt)
            _set_run_font(r, HEAD_FONT, header_font, True, WHITE)

        for ri, row in enumerate(rows):
            tr = t.add_row()
            _cant_split(tr)
            for ci, txt in enumerate(row):
                cell = tr.cells[ci]
                cell.width = Cm(widths[ci])
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                if zebra and ri % 2 == 1:
                    _cell_shade(cell, "F6F8FA")
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.22
                p.paragraph_format.first_line_indent = Pt(0)
                if align and ci < len(align) and align[ci] is not None:
                    p.alignment = align[ci]
                elif widths[ci] < 1.9:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                self._rich(p, txt, size=font,
                           color=NAVY if (ci == 0 and first_col_bold) else INK,
                           bold=(ci == 0 and first_col_bold))

        if len(rows) <= 14:
            for row in t.rows[:-1]:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        para.paragraph_format.keep_with_next = True

        if note:
            n = self.doc.add_paragraph()
            n.paragraph_format.space_before = Pt(2)
            n.paragraph_format.space_after = Pt(10)
            n.paragraph_format.line_spacing = 1.3
            n.paragraph_format.first_line_indent = Pt(0)
            r = n.add_run("注：" + note)
            _set_run_font(r, BODY_FONT, 7.6, False, LIGHTGRAY)
        else:
            self.spacer(8)
