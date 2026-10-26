# -*- coding: utf-8 -*-
"""Assemble the DOCX report, with a two-pass table of contents."""
import os
import re
import sys
import json
import subprocess

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from report_lib import (Builder, setup_document, _set_run_font, _para_border,
                        NAVY, BLUE, INK, GRAY, LIGHTGRAY, AMBER, WHITE,
                        BODY_FONT, HEAD_FONT, CONTENT_WIDTH_CM)
import content_a
import content_b
import content_c

TITLE = "单细胞大模型：问题、方案与架构"
OUT_DOCX = "单细胞大模型_问题与架构_调研报告.docx"
SOFFICE = os.environ.get("SOFFICE") or (
    r"C:\Program Files\LibreOffice\program\soffice.exe"
    if os.name == "nt" else "soffice"
)


# --------------------------------------------------------------------------- probe
class Probe:
    """Records headings without touching a document."""
    def __init__(self):
        self.headings = []

    def h1(self, number, title, pagebreak=True):
        self.headings.append((1, number, title))

    def h2(self, number, title):
        self.headings.append((2, number, title))

    def h3(self, number, title):
        pass

    def __getattr__(self, name):
        def _noop(*args, **kwargs):
            return None
        return _noop


def collect_headings():
    p = Probe()
    content_a.abstract(p)
    content_a.chapter1(p)
    content_a.chapter2(p)
    content_b.chapter3(p)
    content_b.chapter4(p)
    content_c.chapter5(p)
    content_c.chapter6(p)
    content_c.chapter7(p)
    content_c.appendix(p)
    return p.headings


# --------------------------------------------------------------------------- furniture
def _pg_num_type(section, fmt=None, start=None):
    sectPr = section._sectPr
    el = sectPr.find(qn("w:pgNumType"))
    if el is None:
        el = OxmlElement("w:pgNumType")
        sectPr.append(el)
    if fmt:
        el.set(qn("w:fmt"), fmt)
    if start is not None:
        el.set(qn("w:start"), str(start))


def _header_text(paragraph, left, right):
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    r = paragraph.add_run(left)
    _set_run_font(r, HEAD_FONT, 7.8, False, GRAY)
    r = paragraph.add_run("\t" + right)
    _set_run_font(r, HEAD_FONT, 7.8, False, GRAY)
    _para_border(paragraph, "bottom", "D5DCE6", 4, 4)


def _footer_page(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = paragraph.add_run("— ")
    _set_run_font(r, HEAD_FONT, 8.5, False, GRAY)
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "1"
    fld3 = OxmlElement("w:fldChar"); fld3.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
    run._r.append(t); run._r.append(fld3)
    _set_run_font(run, HEAD_FONT, 8.5, False, GRAY)
    r = paragraph.add_run(" —")
    _set_run_font(r, HEAD_FONT, 8.5, False, GRAY)


def build_cover(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True

    def para(before=0, after=0, line=1.0):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = line
        p.paragraph_format.first_line_indent = Pt(0)
        return p

    p = para(before=6)
    _para_border(p, "bottom", "1F3A5F", 16, 2)

    p = para(before=118)
    r = p.add_run("RESEARCH REPORT　·　2026")
    _set_run_font(r, HEAD_FONT, 9, True, AMBER)

    p = para(before=10)
    r = p.add_run("单细胞大模型")
    _set_run_font(r, HEAD_FONT, 40, True, NAVY)

    p = para(before=4, after=14)
    r = p.add_run("问题、方案与架构")
    _set_run_font(r, HEAD_FONT, 40, True, RGBColor(0x2E, 0x5C, 0x8A))

    p = para(before=0, after=16)
    _para_border(p, "bottom", "C0803A", 10, 1)

    p = para(before=0, after=2, line=1.4)
    r = p.add_run("十二个可验证问题 · 六个架构大类 · 二十七个模型的设计取舍")
    _set_run_font(r, HEAD_FONT, 13, False, RGBColor(0x3D, 0x4B, 0x5C))

    p = para(before=4, after=0, line=1.4)
    r = p.add_run("从“把 BERT 搬到细胞”到“用分布预测干预”：一次以代价为中心的技术盘点")
    _set_run_font(r, BODY_FONT, 10, False, GRAY)

    p = para(before=136)
    rows = [
        ("版　　本", "v1.0（汇报版）"),
        ("完成时间", "2026 年 10 月 26 日"),
        ("覆盖范围", "2021–2026 年公开的单细胞基础模型、评测与批评性研究"),
        ("资料口径", "论文原文、预印本、官方仓库与模型卡；厂商口径单独标注"),
        ("引用数量", "27 款模型 · 29 张表 · 47 条来源"),
    ]
    for k, v in rows:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(2.5)
        pf.line_spacing = 1.0
        pf.first_line_indent = Pt(0)
        pf.tab_stops.add_tab_stop(Cm(3.4), WD_TAB_ALIGNMENT.LEFT)
        r = p.add_run(k)
        _set_run_font(r, HEAD_FONT, 8.2, False, LIGHTGRAY)
        r = p.add_run("\t")
        _set_run_font(r, HEAD_FONT, 8.2, False, LIGHTGRAY)
        r = p.add_run(v)
        _set_run_font(r, BODY_FONT, 8.8, False, RGBColor(0x3D, 0x4B, 0x5C))

    p = para(before=14)
    _para_border(p, "top", "D5DCE6", 4, 6)
    p = para(before=4, after=0, line=1.4)
    r = p.add_run("说明：本报告区分“论文/官方数据”与“厂商宣称”，凡属后者均在正文标注；未标注者可在附录 C 找到来源。报告不构成投资建议。")
    _set_run_font(r, BODY_FONT, 7.4, False, LIGHTGRAY)


def build_front(doc, headings, page_map):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    _pg_num_type(sec, "lowerRoman", 1)
    sec.header.is_linked_to_previous = False
    sec.footer.is_linked_to_previous = False
    hdr = sec.header.paragraphs[0]
    _header_text(hdr, TITLE, "摘要与目录")
    _footer_page(sec.footer.paragraphs[0])

    b = Builder(doc, page_map)
    content_a.abstract(b)

    # ---- TOC
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run("目　录")
    _set_run_font(r, HEAD_FONT, 20, True, NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    _para_border(p, "bottom", "C9D2DE", 6, 4)

    for level, number, title in headings:
        if level > 2:
            continue
        anchor = (number + " " + title).strip()
        page = "i" if number.strip() == "摘 要" else page_map.get(anchor, "")
        para = doc.add_paragraph()
        pf = para.paragraph_format
        pf.line_spacing = 1.28
        pf.space_after = Pt(1.5 if level == 2 else 3)
        pf.space_before = Pt(6 if level == 1 else 0)
        pf.left_indent = Cm(0.0 if level == 1 else 0.65)
        pf.first_line_indent = Pt(0)
        pf.tab_stops.add_tab_stop(Cm(CONTENT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        if level == 1:
            label = "摘要" if number.strip() == "摘 要" else f"{number}　{title}"
            r = para.add_run(label)
            _set_run_font(r, HEAD_FONT, 9.6, True, NAVY)
        else:
            r = para.add_run(f"{number}　{title}")
            _set_run_font(r, BODY_FONT, 9.0, False, RGBColor(0x3D, 0x4B, 0x5C))
        r = para.add_run("\t " + str(page))
        _set_run_font(r, HEAD_FONT, 9.0, level == 1, NAVY if level == 1 else GRAY)

    # small note
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run("页码为正文页码；图表编号在全书内连续。")
    _set_run_font(r, BODY_FONT, 7.6, False, LIGHTGRAY)


def build_body(doc, page_map):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    _pg_num_type(sec, "decimal", 1)
    sec.header.is_linked_to_previous = False
    sec.footer.is_linked_to_previous = False
    _header_text(sec.header.paragraphs[0], TITLE, "2026 年 10 月")
    _footer_page(sec.footer.paragraphs[0])

    b = Builder(doc, page_map)
    content_a.chapter1(b)
    content_a.chapter2(b)
    content_b.chapter3(b)
    content_b.chapter4(b)
    content_c.chapter5(b)
    content_c.chapter6(b)
    content_c.chapter7(b)
    content_c.appendix(b)
    return b


def build(page_map=None, path=OUT_DOCX):
    page_map = page_map or {}
    doc = setup_document()
    doc.core_properties.title = "单细胞大模型：问题、方案与架构"
    doc.core_properties.subject = "单细胞基础模型调研报告"
    doc.core_properties.author = "调研报告"
    doc.core_properties.comments = "信息截至 2026-10"

    build_cover(doc)
    headings = collect_headings()
    build_front(doc, headings, page_map)
    body_builder = build_body(doc, page_map)

    doc.save(path)
    return headings, body_builder


def to_pdf(docx_path):
    try:
        subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", ".",
                        docx_path], check=True, timeout=600,
                       stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    except (FileNotFoundError, subprocess.CalledProcessError) as exc:
        sys.exit("导出 PDF 需要 LibreOffice，可用环境变量 SOFFICE 指定 soffice 路径：%s" % exc)
    return os.path.splitext(docx_path)[0] + ".pdf"


def extract_pages(pdf_path, headings):
    import fitz
    doc = fitz.open(pdf_path)
    texts = []
    for page in doc:
        t = page.get_text("text")
        texts.append(re.sub(r"\s+", "", t))
    # body start: page containing chapter 1 lead
    body_start = None
    for i, t in enumerate(texts):
        if "第1章研究范围与分析框架" in t and "在讨论" in t:
            body_start = i
            break
    if body_start is None:
        body_start = 0
    page_map = {}
    for level, number, title in headings:
        if level > 2:
            continue
        key = (number + " " + title).strip()
        needle = re.sub(r"\s+", "", number + title)
        found = None
        for i in range(body_start, len(texts)):
            if needle and needle in texts[i]:
                found = i
                break
        if found is not None:
            page_map[key] = found - body_start + 1
    doc.close()
    return page_map, body_start, len(texts)


def main():
    print("pass 1 ...")
    headings, _ = build({}, OUT_DOCX)
    pdf = to_pdf(OUT_DOCX)
    page_map, body_start, total = extract_pages(pdf, headings)
    print(f"pages={total}, body_start={body_start}, toc_entries={len(page_map)}")
    missing = [h for h in headings if h[0] <= 2 and (h[1] + " " + h[2]).strip() not in page_map]
    print("missing:", missing)

    print("pass 2 ...")
    headings2, bb = build(page_map, OUT_DOCX)
    pdf2 = to_pdf(OUT_DOCX)
    page_map2, body_start2, total2 = extract_pages(pdf2, headings2)
    changed = {k: (page_map.get(k), page_map2.get(k)) for k in page_map2
               if page_map.get(k) != page_map2.get(k)}
    print(f"pass2 pages={total2}, body_start={body_start2}, changed={len(changed)}")
    if changed:
        print("rebuild with pass2 map ...")
        build(page_map2, OUT_DOCX)
        pdf3 = to_pdf(OUT_DOCX)
        print("final pdf:", pdf3)
    print("done")


if __name__ == "__main__":
    main()
